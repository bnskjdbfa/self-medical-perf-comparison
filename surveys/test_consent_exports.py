import io
import zipfile
from datetime import date

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from admins.models import AdminTeamPermission
from teams.models import Respondent, Team

from .consent_exports import (
    build_consent_workbook,
    build_consent_zip,
    render_consent_confirmation_pdf,
)
from .models import (
    ConsentExportAudit,
    DocumentSubmission,
    DocumentTemplate,
    RequiredDocumentConsentBundle,
)


class ConsentExportTest(TestCase):
    def setUp(self):
        self.team = Team.objects.create(
            name='1차 원정대',
            organization_name='운영기관',
            organization_name_en='Operating Organization',
            code='CONSENT-EXPORT',
            is_active=True,
            start_date=date.today(),
            max_day_no=3,
        )
        self.respondent = Respondent.objects.create(
            team=self.team,
            name='현재 이름',
            birth_date=date(1990, 1, 2),
            phone='01000000000',
            gender='M',
            normalized_name='현재이름',
            normalized_phone='01000000000',
        )
        self.bundle = RequiredDocumentConsentBundle.objects.create(
            team=self.team,
            respondent=self.respondent,
            participant_snapshot_json={
                'name': '=제출당시이름',
                'birth_date': '1990-01-02',
                'phone': '01012345678',
                'gender': 'F',
            },
            team_snapshot_json={
                'team_name': '제출 당시 팀',
                'organization_name': '제출 당시 기관',
            },
            enabled_document_types_snapshot=[
                DocumentTemplate.DocumentType.PRIVACY_CONSENT,
                DocumentTemplate.DocumentType.THIRD_PARTY_CONSENT,
                DocumentTemplate.DocumentType.SAFETY_PLEDGE,
            ],
            completed_at=timezone.now(),
        )
        self.submission = DocumentSubmission.objects.create(
            team=self.team,
            respondent=self.respondent,
            bundle=self.bundle,
            document_type=DocumentTemplate.DocumentType.PRIVACY_CONSENT,
            template_version=2,
            template_hash_at_submit='a' * 64,
            document_title_snapshot='개인정보 수집·이용 동의서',
            document_content_snapshot='<p><strong>제출 당시 본문</strong><br>두 번째 줄</p>',
            document_content_hash_at_submit='b' * 64,
            agreed=True,
            signer_name='제출 당시 이름',
            signature_method=DocumentSubmission.SignatureMethod.CHECKBOX,
            agreed_at=timezone.now(),
        )

    def test_pdf_uses_submission_and_participant_snapshots(self):
        pdf = render_consent_confirmation_pdf(
            DocumentSubmission.objects.select_related('team', 'respondent', 'bundle').get(
                pk=self.submission.pk
            )
        )

        self.assertTrue(pdf.startswith(b'%PDF-'))
        self.assertGreater(len(pdf), 3000)

    def test_excel_has_two_sheets_and_escapes_formula_values(self):
        bundle = (
            RequiredDocumentConsentBundle.objects.select_related('team', 'respondent')
            .prefetch_related('submissions')
            .get(pk=self.bundle.pk)
        )
        content = build_consent_workbook([bundle])
        workbook = load_workbook(io.BytesIO(content))

        self.assertEqual(workbook.sheetnames, ['참가자별 제출 현황', '문서별 제출 상세'])
        status = workbook['참가자별 제출 현황']
        self.assertEqual(status['A2'].value, '제출 당시 기관')
        self.assertEqual(status['B2'].value, '제출 당시 팀')
        self.assertEqual(status['C2'].value, "'=제출당시이름")
        self.assertEqual(status['E2'].value, '01012345678')
        self.assertEqual(status['F2'].value, '여')
        self.assertEqual(status['G2'].value, '제출 완료')
        self.assertEqual(status['I2'].value, '미제출')
        self.assertEqual(status['M2'].value, '완료')
        detail = workbook['문서별 제출 상세']
        self.assertEqual(detail['C2'].value, "'=제출당시이름")
        self.assertEqual(detail['E2'].value, 2)
        self.assertEqual(detail['I2'].value, 'b' * 64)

    def test_zip_contains_individual_confirmation_pdf(self):
        archive = build_consent_zip(
            DocumentSubmission.objects.select_related('team', 'respondent', 'bundle')
        )

        with zipfile.ZipFile(io.BytesIO(archive.content)) as zipped:
            self.assertEqual(len(zipped.namelist()), 1)
            self.assertTrue(zipped.read(zipped.namelist()[0]).startswith(b'%PDF-'))
        self.assertEqual(archive.submission_ids, (self.submission.pk,))

    def test_team_admin_downloads_excel_and_zip_and_records_audit(self):
        user = get_user_model().objects.create_superuser(
            username='consent-admin', password='pass', email='admin@example.com'
        )
        self.client.force_login(user)

        excel = self.client.get(reverse('admin:teams_team_consent_excel', args=[self.team.pk]))
        zipped = self.client.get(reverse('admin:teams_team_consent_zip', args=[self.team.pk]))

        self.assertEqual(excel.status_code, 200)
        self.assertEqual(
            excel['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertEqual(zipped.status_code, 200)
        self.assertEqual(zipped['Content-Type'], 'application/zip')
        self.assertEqual(ConsentExportAudit.objects.filter(result='success').count(), 2)

    def test_individual_pdf_endpoint_reapplies_team_permission(self):
        user_model = get_user_model()
        allowed = user_model.objects.create_user(
            username='consent-allowed', password='pass', is_staff=True
        )
        allowed.user_permissions.add(Permission.objects.get(codename='view_documentsubmission'))
        AdminTeamPermission.objects.create(user=allowed, team=self.team)
        denied = user_model.objects.create_user(
            username='consent-denied', password='pass', is_staff=True
        )
        denied.user_permissions.add(Permission.objects.get(codename='view_documentsubmission'))
        url = reverse(
            'admin:surveys_documentsubmission_confirmation_pdf',
            args=[self.submission.pk],
        )

        self.client.force_login(denied)
        self.assertEqual(self.client.get(url).status_code, 302)
        # Exercise the real login/key rotation rather than bypassing generation
        # credentials with TestClient.force_login on an already active browser.
        self.client.post(reverse('admin:login'), {'username': allowed.username, 'password': 'pass'})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF-'))
        audit = ConsentExportAudit.objects.get(export_format='pdf')
        self.assertEqual(audit.admin_user, allowed)
        self.assertEqual(audit.team, self.team)
