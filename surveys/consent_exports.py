from __future__ import annotations

import html
import io
import re
import zipfile
from dataclasses import dataclass

from django.utils import timezone
from django.utils.html import strip_tags
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from teams.models import GenderChoices

from .models import DocumentSubmission, DocumentTemplate, RequiredDocumentConsentBundle
from .reflection_exports import safe_filename


EXCEL_FORMULA_PREFIXES = ('=', '+', '-', '@')
DOCUMENT_LABELS = dict(DocumentTemplate.DocumentType.choices)


class ConsentExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class ConsentZipArchive:
    content: bytes
    file_name: str
    submission_ids: tuple[int, ...]


def safe_excel_text(value) -> str:
    text = str(value or '')
    return f"'{text}" if text.startswith(EXCEL_FORMULA_PREFIXES) else text


def _participant_snapshot(bundle: RequiredDocumentConsentBundle) -> dict:
    snapshot = bundle.participant_snapshot_json or {}
    return {
        'name': snapshot.get('name') or bundle.respondent.name,
        'birth_date': snapshot.get('birth_date') or bundle.respondent.birth_date.isoformat(),
        'phone': snapshot.get('phone') or bundle.respondent.phone,
        'gender': snapshot.get('gender') or bundle.respondent.gender or '',
    }


def _team_snapshot(bundle: RequiredDocumentConsentBundle) -> dict:
    snapshot = bundle.team_snapshot_json or {}
    return {
        'organization_name': snapshot.get('organization_name') or bundle.team.organization_name,
        'team_name': snapshot.get('team_name') or bundle.team.name,
    }


def _local_text(value) -> str:
    if not value:
        return ''
    return timezone.localtime(value).strftime('%Y-%m-%d %H:%M:%S')


def consent_excel_filename(bundles: list[RequiredDocumentConsentBundle]) -> str:
    if not bundles:
        return f'동의서_제출현황_{timezone.localdate().isoformat()}.xlsx'
    snapshots = [_team_snapshot(bundle) for bundle in bundles]
    organizations = {item['organization_name'] for item in snapshots}
    teams = {item['team_name'] for item in snapshots}
    prefix = '_'.join(
        filter(
            None,
            [
                next(iter(organizations)) if len(organizations) == 1 else '복수기관',
                next(iter(teams)) if len(teams) == 1 else '복수팀',
            ],
        )
    )
    return f'{safe_filename(prefix, fallback="동의서")}_동의서제출현황_{timezone.localdate().isoformat()}.xlsx'


def build_consent_workbook(bundles) -> bytes:
    bundles = list(bundles)
    if not bundles:
        raise ConsentExportError('내보낼 동의서 제출 내역이 없습니다.')

    workbook = Workbook()
    status_sheet = workbook.active
    status_sheet.title = '참가자별 제출 현황'
    detail_sheet = workbook.create_sheet('문서별 제출 상세')
    status_headers = [
        '운영기관', '팀', '이름', '생년월일', '휴대전화', '성별',
        '개인정보 동의 상태', '개인정보 동의 일시',
        '제3자 제공 동의 상태', '제3자 제공 동의 일시',
        '안전서약 상태', '안전서약 일시', '전체 완료 여부',
    ]
    detail_headers = [
        '운영기관', '팀', '참가자', '문서 유형', '문서 버전',
        '동의 여부', '제출 일시', '문서 제목', '본문 SHA-256',
    ]
    status_sheet.append(status_headers)
    detail_sheet.append(detail_headers)

    for bundle in bundles:
        participant = _participant_snapshot(bundle)
        team = _team_snapshot(bundle)
        prefetched = getattr(bundle, 'export_submissions', None)
        bundle_submissions = prefetched if prefetched is not None else bundle.submissions.all()
        submissions = {
            submission.document_type: submission
            for submission in bundle_submissions
        }
        values = []
        for document_type in (
            DocumentTemplate.DocumentType.PRIVACY_CONSENT,
            DocumentTemplate.DocumentType.THIRD_PARTY_CONSENT,
            DocumentTemplate.DocumentType.SAFETY_PLEDGE,
        ):
            submission = submissions.get(document_type)
            if document_type not in bundle.enabled_document_types_snapshot:
                values.extend(['미사용', ''])
            elif submission:
                values.extend(['제출 완료', _local_text(submission.agreed_at)])
            else:
                values.extend(['미제출', ''])
        gender = dict(GenderChoices.choices).get(participant['gender'], participant['gender'])
        status_sheet.append([
            safe_excel_text(team['organization_name']), safe_excel_text(team['team_name']),
            safe_excel_text(participant['name']), participant['birth_date'],
            safe_excel_text(participant['phone']), safe_excel_text(gender), *values,
            '완료' if bundle.completed_at else '미완료',
        ])
        for submission in submissions.values():
            detail_sheet.append([
                safe_excel_text(team['organization_name']), safe_excel_text(team['team_name']),
                safe_excel_text(participant['name']),
                DOCUMENT_LABELS.get(submission.document_type, submission.document_type),
                submission.template_version, '동의' if submission.agreed else '미동의',
                _local_text(submission.agreed_at),
                safe_excel_text(submission.document_title_snapshot),
                submission.document_content_hash_at_submit,
            ])

    for sheet in (status_sheet, detail_sheet):
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='477F99')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for column in range(1, sheet.max_column + 1):
            width = max(len(str(sheet.cell(row, column).value or '')) for row in range(1, sheet.max_row + 1))
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 11), 42)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_fonts() -> tuple[str, str]:
    font = 'HYSMyeongJo-Medium'
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    return font, font


def _paragraph_text(value) -> str:
    return html.escape(str(value or '')).replace('\n', '<br/>')


def _document_plain_text(content: str) -> str:
    with_breaks = re.sub(r'(?i)<br\s*/?>|</p>|</li>|</h[1-6]>', '\n', content or '')
    return html.unescape(strip_tags(with_breaks)).strip()


def render_consent_confirmation_pdf(submission: DocumentSubmission) -> bytes:
    if not submission.document_content_snapshot:
        raise ConsentExportError('제출 당시 문서 본문을 복원할 수 없습니다.')
    bundle = submission.bundle
    if bundle is None:
        raise ConsentExportError('제출 당시 참가자 스냅샷을 복원할 수 없습니다.')
    participant = _participant_snapshot(bundle)
    team = _team_snapshot(bundle)
    gender = dict(GenderChoices.choices).get(participant['gender'], participant['gender'])
    regular, bold = _pdf_fonts()
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title=f'{team["team_name"]}_{participant["name"]}_전자 동의 확인서',
        author=team['organization_name'],
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('ConsentTitle', parent=styles['Title'], fontName=bold, fontSize=19, leading=26, alignment=TA_CENTER)
    normal = ParagraphStyle('ConsentNormal', parent=styles['Normal'], fontName=regular, fontSize=9.5, leading=15)
    heading = ParagraphStyle('ConsentHeading', parent=normal, fontName=bold, fontSize=12, leading=18, spaceBefore=5 * mm)
    metadata = [
        ['운영기관', team['organization_name'], '팀', team['team_name']],
        ['이름', participant['name'], '생년월일', participant['birth_date']],
        ['휴대전화', participant['phone'], '성별', gender],
        ['문서 유형', DOCUMENT_LABELS.get(submission.document_type, submission.document_type), '버전', str(submission.template_version)],
        ['동의 결과', '동의' if submission.agreed else '미동의', '동의 일시', _local_text(submission.agreed_at)],
        ['동의 방식', submission.get_signature_method_display(), '본문 hash', submission.document_content_hash_at_submit],
    ]
    table = Table(
        [[Paragraph(_paragraph_text(value), normal) for value in row] for row in metadata],
        colWidths=[25 * mm, 60 * mm, 25 * mm, 60 * mm],
    )
    table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#BBBBBB')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F1F1F1')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#F1F1F1')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story = [
        Paragraph('전자 동의 확인서', title_style), Spacer(1, 7 * mm), table,
        Paragraph(_paragraph_text(submission.document_title_snapshot), heading),
    ]
    for paragraph in _document_plain_text(submission.document_content_snapshot).split('\n'):
        if paragraph.strip():
            story.append(Paragraph(_paragraph_text(paragraph.strip()), normal))
            story.append(Spacer(1, 1.5 * mm))
    story.extend([
        PageBreak(),
        Paragraph('전자 동의 증빙정보', heading),
        Paragraph(f'문서 본문 SHA-256: {_paragraph_text(submission.document_content_hash_at_submit)}', normal),
        Paragraph(f'제출 당시 문서 버전: {submission.template_version}', normal),
        Paragraph(f'동의 결과: {"동의" if submission.agreed else "미동의"}', normal),
        Paragraph(f'동의 일시: {_local_text(submission.agreed_at)}', normal),
    ])
    document.build(story)
    return output.getvalue()


def consent_pdf_filename(submission: DocumentSubmission) -> str:
    bundle = submission.bundle
    if bundle is None:
        raise ConsentExportError('제출 당시 참가자 스냅샷을 복원할 수 없습니다.')
    participant = _participant_snapshot(bundle)
    team = _team_snapshot(bundle)
    label = DOCUMENT_LABELS.get(submission.document_type, submission.document_type)
    stem = '_'.join((team['organization_name'], team['team_name'], participant['name'], label, '전자동의확인서'))
    return f'{safe_filename(stem, fallback=f"consent-{submission.pk}")}.pdf'


def build_consent_zip(submissions, *, file_prefix: str = '동의서') -> ConsentZipArchive:
    submissions = list(submissions)
    if not submissions:
        raise ConsentExportError('내보낼 동의서 제출 내역이 없습니다.')
    output = io.BytesIO()
    used_names: set[str] = set()
    with zipfile.ZipFile(output, mode='w', compression=zipfile.ZIP_DEFLATED) as archive:
        for submission in submissions:
            pdf_content = render_consent_confirmation_pdf(submission)
            name = consent_pdf_filename(submission)
            if name in used_names:
                stem = name[:-4]
                suffix = 2
                while f'{stem}_{suffix}.pdf' in used_names:
                    suffix += 1
                name = f'{stem}_{suffix}.pdf'
            used_names.add(name)
            archive.writestr(name, pdf_content)
    archive_name = f'{safe_filename(file_prefix, fallback="동의서")}_전자동의확인서_{timezone.localdate().isoformat()}.zip'
    return ConsentZipArchive(output.getvalue(), archive_name, tuple(item.pk for item in submissions))
