from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import AnonymousUser
from django.db import connection
from django.test import Client, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from admins.models import AdminTeamPermission
from admins.permissions import allowed_team_ids
from surveys.models import (
    AdminDownloadAudit,
    ExitSurveySubmission,
    ReflectionSubmission,
    RequiredDocumentConsentBundle,
    SurveyAnswer,
    SurveyForm,
    SurveyQuestion,
    SurveySection,
    SurveySubmission,
)
from teams.models import Participant, Respondent, Team


class AdminPagesTest(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.staff = user_model.objects.create_user(username='staff', password='pass', is_staff=True)
        self.user = user_model.objects.create_user(username='user', password='pass', is_staff=False)

        self.team = Team.objects.create(name='팀A', code='1234', manager_entry_code='team-a-admin', is_active=True, start_date=date.today(), max_day_no=3)
        self.other_team = Team.objects.create(name='팀B', code='5678', manager_entry_code='team-b-admin', is_active=True, start_date=date.today(), max_day_no=3)
        respondent = Respondent.objects.create(
            team=self.team,
            name='홍길동',
            birth_date='1985-03-12',
            phone='01012345678',
            gender='M',
            normalized_name='홍길동',
            normalized_phone='01012345678',
        )
        Participant.objects.create(
            team=self.team,
            name='홍길동',
            birth_date='1985-03-12',
            phone='01012345678',
            gender='M',
            normalized_name='홍길동',
            normalized_phone='01012345678',
            is_active=True,
        )
        Participant.objects.create(
            team=self.team,
            name='김철수',
            birth_date='1990-07-01',
            phone='01022223333',
            normalized_name='김철수',
            normalized_phone='01022223333',
            is_active=True,
        )
        form = SurveyForm.objects.create(team=self.team, day_no=1, title='1일차', version=1, schema_json={}, is_active=True)
        section = SurveySection.objects.create(form=form, title='신체', section_type='physical', display_order=1)
        self.question = SurveyQuestion.objects.create(
            section=section,
            key='q1',
            label='두통 여부',
            question_type='single_choice',
            options_json=['예', '아니오'],
            is_required=True,
            display_order=1,
        )
        self.text_question = SurveyQuestion.objects.create(
            section=section,
            key='q_note',
            label='추가로 느낀 증상을 적어주세요',
            question_type='text',
            options_json=[],
            is_required=False,
            display_order=2,
        )
        submission = SurveySubmission.objects.create(
            team=self.team,
            respondent=respondent,
            day_no=1,
            answers_json={'q1': '예', 'q_note': '두통과 어지러움이 있습니다.'},
            result_json={
                'status': 'submitted',
                'sections': {
                    'body': {'average_score': 25, 'status_label': '경고', 'surveys': []},
                    'emotion': {'average_score': 72, 'status_label': '정상', 'surveys': []},
                },
            },
            form_version=1,
            form_snapshot_json={},
            is_locked=True,
            submitted_at=timezone.now(),
        )
        SurveyAnswer.objects.create(
            submission=submission,
            question=self.question,
            answer_text='예',
            answer_json={'value': '예'},
        )
        SurveyAnswer.objects.create(
            submission=submission,
            question=self.text_question,
            answer_text='두통과 어지러움이 있습니다.',
            answer_json={'value': '두통과 어지러움이 있습니다.'},
        )
        AdminTeamPermission.objects.create(user=self.staff, team=self.team, role=AdminTeamPermission.Roles.READ_ONLY_TEAM_ADMIN)

        other_respondent = Respondent.objects.create(
            team=self.other_team,
            name='이영희',
            birth_date='1988-02-20',
            phone='01099998888',
            normalized_name='이영희',
            normalized_phone='01099998888',
        )
        other_submission = SurveySubmission.objects.create(
            team=self.other_team,
            respondent=other_respondent,
            day_no=1,
            answers_json={'q1': '예'},
            result_json={
                'status': 'submitted',
                'sections': {
                    'body': {'average_score': 90, 'status_label': '최상', 'surveys': []},
                    'emotion': {'average_score': 84, 'status_label': '최상', 'surveys': []},
                },
            },
            form_version=1,
            form_snapshot_json={},
            is_locked=True,
            submitted_at=timezone.now(),
        )
        self.submission_id = submission.id
        self.other_submission_id = other_submission.id

    def test_admin_login_page_renders_for_anonymous_user(self):
        response = self.client.get(reverse('admin:login'))

        self.assertEqual(response.status_code, 200)

    def test_anonymous_user_has_no_allowed_teams(self):
        self.assertEqual(allowed_team_ids(AnonymousUser()), [])

    def test_staff_can_access_dashboard(self):
        self.client.login(username='staff', password='pass')
        response = self.client.get(reverse('admins:dashboard'))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])

    def test_non_staff_forbidden(self):
        self.client.login(username='user', password='pass')
        response = self.client.get(reverse('admins:dashboard'))
        self.assertEqual(response.status_code, 302)

    def test_staff_can_view_response_detail(self):
        self.client.login(username='staff', password='pass')
        response = self.client.get(reverse('admins:response-detail', kwargs={'submission_id': self.submission_id}))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['data']['answers'][0]['question_label'], '두통 여부')
        self.assertEqual(
            response.json()['data']['subjective_answers'][0]['question_label'],
            '추가로 느낀 증상을 적어주세요',
        )

    def test_response_detail_uses_submission_snapshot_without_answer_rows(self):
        submission = SurveySubmission.objects.get(pk=self.submission_id)
        submission.form_snapshot_json = {
            'sections': [
                {
                    'title': '제출 당시 신체 섹션',
                    'questions': [
                        {
                            'id': self.question.id,
                            'key': 'q1',
                            'label': '제출 당시 두통 여부',
                            'question_type': 'single',
                        },
                        {
                            'id': self.text_question.id,
                            'key': 'q_note',
                            'label': '제출 당시 추가 증상',
                            'question_type': 'text',
                        },
                    ],
                },
            ],
        }
        submission.save(update_fields=['form_snapshot_json', 'updated_at'])
        submission.answers.all().delete()
        self.client.login(username='staff', password='pass')

        response = self.client.get(
            reverse('admins:response-detail', kwargs={'submission_id': self.submission_id})
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()['data']
        self.assertEqual(
            [answer['question_label'] for answer in data['answers']],
            ['제출 당시 두통 여부', '제출 당시 추가 증상'],
        )
        self.assertEqual(data['answers'][0]['display_answer'], '예')
        self.assertEqual(data['subjective_answers'][0]['display_answer'], '두통과 어지러움이 있습니다.')

    def test_staff_can_access_manager_prototype(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '미제출')
        self.assertNotContains(response, self.other_team.name)
        self.assertContains(response, '대상자 수')
        self.assertContains(response, '홍길동')
        html = response.content.decode()
        self.assertNotIn('{%', html)
        self.assertNotIn('%}', html)

    def test_manager_matching_treats_gender_mismatch_as_pending_when_both_exist(self):
        participant = Participant.objects.get(team=self.team, normalized_name='홍길동')
        participant.gender = 'F'
        participant.save(update_fields=['gender', 'updated_at'])

        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 200)
        rows = response.context['rows']
        self.assertEqual(rows[0]['status_label'], '미제출')

    def test_superuser_can_access_admin_site(self):
        user_model = get_user_model()
        user_model.objects.create_superuser(username='root', password='pass', email='root@example.com')
        self.client.login(username='root', password='pass')
        response = self.client.get('/admin/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '동의서 관리')
        self.assertContains(response, '설문 및 검진 관리')
        self.assertNotContains(response, '>검진 관리<', html=False)
        self.assertNotContains(response, '>설문 관리<', html=False)
        self.assertContains(response, '소감문 관리')

        grouped_apps = {
            app['name']: {
                model['object_name']
                for model in app['models']
            }
            for app in response.context['available_apps']
        }
        self.assertEqual(
            grouped_apps['동의서 관리'],
            {'DocumentSubmission', 'RequiredDocumentConsentBundle'},
        )
        self.assertEqual(
            grouped_apps['소감문 관리'],
            {'ReflectionSubmission'},
        )
        combined_group = grouped_apps['설문 및 검진 관리']
        self.assertIn('SurveySubmission', combined_group)
        self.assertIn('CommonQuestionSet', combined_group)
        self.assertIn('ExitSurveySubmission', combined_group)
        combined_models = {
            model['object_name']: model['name']
            for app in response.context['available_apps']
            if app['name'] == '설문 및 검진 관리'
            for model in app['models']
        }
        self.assertEqual(combined_models['SurveySubmission'], '검진 제출')
        self.assertEqual(combined_models['ExitSurveySubmission'], '설문 제출')
        visible_models = {
            model_name
            for model_names in grouped_apps.values()
            for model_name in model_names
        }
        self.assertNotIn('SurveyForm', visible_models)
        self.assertNotIn('SurveySection', visible_models)
        self.assertNotIn('SurveyQuestion', visible_models)
        self.assertNotIn('SurveyAnswer', visible_models)
        self.assertNotIn('TeamCommonQuestionExclusion', visible_models)

    def test_superuser_can_open_survey_admin_changelists(self):
        user_model = get_user_model()
        user_model.objects.create_superuser(username='root2', password='pass', email='root2@example.com')
        self.client.login(username='root2', password='pass')

        self.assertEqual(self.client.get(reverse('admin:surveys_surveysection_changelist')).status_code, 200)
        self.assertEqual(self.client.get(reverse('admin:surveys_surveyquestion_changelist')).status_code, 200)
        self.assertEqual(self.client.get(reverse('admin:surveys_surveyanswer_changelist')).status_code, 200)

    def test_superuser_can_open_health_submission_detail_with_answer_overview(self):
        user_model = get_user_model()
        user_model.objects.create_superuser(
            username='root-health-detail',
            password='pass',
            email='root-health-detail@example.com',
        )
        self.client.login(username='root-health-detail', password='pass')

        response = self.client.get(
            reverse(
                'admin:surveys_surveysubmission_change',
                args=(self.submission_id,),
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '답변 요약')
        self.assertContains(response, '두통 여부')
        self.assertNotContains(response, 'Result json')

    def test_manager_default_day_uses_team_start_date_based_day(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['filters']['day_no'], 1)

    def _login_verified_manager(self):
        root = get_user_model().objects.create_superuser(
            username='verified-manager', password='pass', email='verified-manager@example.com'
        )
        self.client.force_login(root)
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

    def test_manager_lookup_renders_health_tab_for_verified_manager(self):
        self._login_verified_manager()
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.team.name)
        self.assertContains(response, 'Research filter')
        self.assertContains(response, 'tel:01012345678')
        self.assertEqual(response.context['active_tab'], 'health')
        self.assertContains(response, 'manager-result-score')

    def test_manager_non_health_tabs_hide_health_only_filters_and_scores(self):
        self._login_verified_manager()
        for tab in ('documents', 'reflection', 'survey'):
            response = self.client.get(reverse('admins:prototype-manager'), {'tab': tab})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.context['active_tab'], tab)
            self.assertNotContains(response, 'name="day_no"')
            self.assertNotContains(response, 'name="risk_level"')
            self.assertNotContains(response, 'manager-result-score')
            self.assertContains(response, 'name="submission_status"')
            self.assertContains(response, 'name="name_keyword"')

    def test_manager_status_tabs_use_their_own_submission_models(self):
        self._login_verified_manager()
        respondent = Respondent.objects.get(team=self.team, normalized_name='홍길동')
        RequiredDocumentConsentBundle.objects.create(
            team=self.team,
            respondent=respondent,
            completed_at=timezone.now(),
        )
        ReflectionSubmission.objects.create(
            team=self.team,
            respondent=respondent,
            content='테스트 소감문',
            submitted_at=timezone.now(),
        )
        exit_form = SurveyForm.objects.create(
            team=self.team,
            day_no=1,
            title='종료 설문',
            version=2,
            schema_json={},
            is_active=True,
            form_type=SurveyForm.FormType.EXIT_SURVEY,
        )
        ExitSurveySubmission.objects.create(
            team=self.team,
            respondent=respondent,
            survey_form=exit_form,
            form_version=2,
            form_snapshot_json={},
            answers_json={},
            submitted_at=timezone.now(),
        )

        for tab in ('documents', 'reflection', 'survey'):
            response = self.client.get(
                reverse('admins:prototype-manager'),
                {'tab': tab, 'submission_status': 'submitted'},
            )
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.context['summary']['submitted_count'], 1)
            self.assertContains(response, '홍길동')
            self.assertNotContains(response, '김철수')

    def test_manager_risk_filter_uses_submission_result_sections(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()
        response = self.client.get(reverse('admins:prototype-manager'), {'risk_level': '경고'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '홍길동')
        self.assertNotContains(response, '김철수')
        self.assertEqual(response.context['summary']['danger_count'], 1)

    def test_manager_list_paginates_filtered_rows(self):
        for index in range(25):
            Participant.objects.create(
                team=self.team,
                name=f'테스트{index:02d}',
                birth_date=date(1991, 1, (index % 28) + 1),
                phone=f'0109000{index:04d}',
                normalized_name=f'테스트{index:02d}',
                normalized_phone=f'0109000{index:04d}',
                is_active=True,
            )

        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

        first_page = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(first_page.status_code, 200)
        self.assertEqual(first_page.context['pagination']['page_size'], 20)
        self.assertEqual(first_page.context['pagination']['page'], 1)
        self.assertEqual(len(first_page.context['rows']), 20)
        self.assertEqual(first_page.context['pagination']['total_count'], 27)
        self.assertTrue(first_page.context['pagination']['has_next'])
        self.assertContains(first_page, '대상자 목록 페이지')
        self.assertContains(first_page, 'page=2', html=False)

        second_page = self.client.get(reverse('admins:prototype-manager'), {'page': 2})
        self.assertEqual(second_page.status_code, 200)
        self.assertEqual(second_page.context['pagination']['page'], 2)
        self.assertEqual(len(second_page.context['rows']), 7)
        self.assertContains(second_page, '테스트24')

    def test_manager_lookup_query_count_stays_bounded_with_more_participants(self):
        for index in range(30):
            Participant.objects.create(
                team=self.team,
                name=f'부하{index:02d}',
                birth_date=date(1992, 1, (index % 28) + 1),
                phone=f'0108000{index:04d}',
                normalized_name=f'부하{index:02d}',
                normalized_phone=f'0108000{index:04d}',
                is_active=True,
            )

        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

        # Establish the one-time legacy session generation outside the steady
        # state lookup budget; pagination must still stay independent of size.
        self.client.get(reverse('admins:prototype-manager'))
        with CaptureQueriesContext(connection) as captured:
            response = self.client.get(reverse('admins:prototype-manager'))

        self.assertEqual(response.status_code, 200)
        # PostgreSQL adds transaction-local safety settings, not a lookup
        # query per participant. Keep the existing lookup budget for all other SQL.
        lookup_queries = [q for q in captured
                          if not q['sql'].startswith("SELECT set_config('lock_timeout',")]
        self.assertLessEqual(len(lookup_queries), 10)
        self.assertEqual(response.context['pagination']['total_count'], 32)
        self.assertEqual(len(response.context['rows']), 20)

    def test_manager_lookup_filter_options_render_selected_values(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()
        response = self.client.get(
            reverse('admins:prototype-manager'),
            {'day_no': 2, 'submission_status': 'pending', 'risk_level': '경고'},
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertNotIn('{%', html)
        self.assertNotIn('%}', html)
        self.assertRegex(html, r'<option\s+value="2"\s+selected\s*>')
        self.assertRegex(html, r'<option\s+value="pending"\s+selected\s*>')
        self.assertRegex(html, r'<option\s+value="경고"\s+selected\s*>')

    def test_submitted_manager_row_contains_result_url_with_return_to(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()
        response = self.client.get(
            reverse('admins:prototype-manager'),
            {'day_no': 1, 'submission_status': 'all', 'risk_level': 'all', 'name_keyword': '홍'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            reverse('surveys:result', kwargs={'day_no': 1}),
            html=False,
        )
        self.assertContains(
            response,
            'source=manager',
            html=False,
        )
        self.assertContains(
            response,
            f'submission_id={self.submission_id}',
            html=False,
        )
        self.assertContains(
            response,
            'return_to=%2Fops%2Fprototype%2Fmanager%2F%3Fday_no%3D1',
            html=False,
        )

    def test_manager_lookup_uses_selected_day_submission_consistently(self):
        respondent = Respondent.objects.get(team=self.team, normalized_name='홍길동')
        day2_submission = SurveySubmission.objects.create(
            team=self.team,
            respondent=respondent,
            day_no=2,
            answers_json={'q1': '아니오'},
            result_json={
                'status': 'submitted',
                'sections': {
                    'body': {'average_score': 92, 'status_label': '최상', 'surveys': []},
                    'emotion': {'average_score': 88, 'status_label': '최상', 'surveys': []},
                },
            },
            form_version=1,
            form_snapshot_json={},
            is_locked=True,
            submitted_at=timezone.now(),
        )
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

        response = self.client.get(
            reverse('admins:prototype-manager'),
            {'day_no': 2, 'submission_status': 'submitted', 'risk_level': 'all', 'name_keyword': '홍'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['rows'][0]['result_url'].split('?')[0], reverse('surveys:result', kwargs={'day_no': 2}))
        self.assertIn(f'submission_id={day2_submission.id}', response.context['rows'][0]['result_url'])
        self.assertContains(response, '최상')

    def test_manager_can_open_result_without_user_survey_session(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

        response = self.client.get(
            reverse('surveys:result', kwargs={'day_no': 1}),
            {
                'source': 'manager',
                'submission_id': self.submission_id,
                'return_to': reverse('admins:prototype-manager'),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '홍길동')
        self.assertContains(
            response,
            f'submission_id={self.submission_id}',
            html=False,
        )

    def test_manager_result_day_tabs_keep_manager_context_without_404(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

        response = self.client.get(
            reverse('surveys:result', kwargs={'day_no': 2}),
            {
                'source': 'manager',
                'submission_id': self.submission_id,
                'return_to': reverse('admins:prototype-manager') + '?day_no=1',
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '홍길동')
        self.assertContains(
            response,
            '/surveys/day/3/result/?source=manager&amp;return_to=%2Fops%2Fprototype%2Fmanager%2F%3Fday_no%3D1',
            html=False,
        )
        self.assertContains(
            response,
            f'submission_id={self.submission_id}',
            html=False,
        )

    def test_manager_result_blocks_submission_from_other_team(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session['ops_login_verified_team_id'] = self.team.id
        session.save()

        response = self.client.get(
            reverse('surveys:result', kwargs={'day_no': 1}),
            {'source': 'manager', 'submission_id': self.other_submission_id},
        )
        self.assertEqual(response.status_code, 404)

    def test_response_detail_api_blocks_not_allowed_team(self):
        self.client.login(username='staff', password='pass')
        response = self.client.get(reverse('admins:response-detail', kwargs={'submission_id': self.other_submission_id}))
        self.assertEqual(response.status_code, 403)

    def test_manager_login_page_renders(self):
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session.save()
        response = self.client.get(reverse('admins:login'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '탐방 통합 플랫폼')
        self.assertContains(response, 'Continue')
        self.assertContains(response, 'method="post"')
        self.assertContains(response, 'name="username"')
        self.assertContains(response, 'name="password"')
        self.assertContains(response, 'name="team_id"')
        self.assertContains(response, 'csrfmiddlewaretoken')
        self.assertContains(response, 'admins/icons/v2-login-key.svg')
        self.assertContains(response, 'v2-primary-button manager-login-cta')

    def test_manager_login_post_remains_csrf_protected(self):
        csrf_client = Client(enforce_csrf_checks=True)
        session = csrf_client.session
        session['ops_login_team_id'] = self.team.id
        session.save()
        response = csrf_client.post(reverse('admins:login'), {'username': 'staff', 'password': 'pass'})
        self.assertEqual(response.status_code, 403)

    def test_manager_login_blocks_user_without_scoped_team_permission(self):
        session = self.client.session
        session['ops_login_team_id'] = self.other_team.id
        session.save()
        response = self.client.post(reverse('admins:login'), {'username': 'staff', 'password': 'pass'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '해당 팀 조회 권한이 없습니다.')

    def test_manager_login_success_redirects_to_manager_page(self):
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session.save()
        response = self.client.post(reverse('admins:login'), {'username': 'staff', 'password': 'pass'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('admins:prototype-manager'))
        self.assertEqual(self.client.session.get('ops_login_verified_team_id'), self.team.id)

    def test_manager_login_shows_form_even_if_user_is_already_authenticated(self):
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session.save()
        self.client.login(username='staff', password='pass')
        response = self.client.get(reverse('admins:login'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Continue')

    def test_prototype_manager_redirects_unauthenticated_user_to_custom_manager_login(self):
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session.save()
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('admins:login'))

    def test_prototype_manager_redirects_to_v2_code_entry_without_scoped_team(self):
        self.client.login(username='staff', password='pass')
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('surveys:preview-v2-code'))

    def test_manager_preview_url_is_removed(self):
        self.assertEqual(self.client.get('/ops/prototype/manager-preview/').status_code, 404)

    def test_prototype_manager_requires_verified_team_session_even_when_authenticated(self):
        self.client.login(username='staff', password='pass')
        session = self.client.session
        session['ops_login_team_id'] = self.team.id
        session.pop('ops_login_verified_team_id', None)
        session.save()
        response = self.client.get(reverse('admins:prototype-manager'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse('admins:login'))

    def test_superuser_can_see_answer_context_in_answer_admin(self):
        user_model = get_user_model()
        user_model.objects.create_superuser(username='root-answers', password='pass', email='root-answers@example.com')
        self.client.login(username='root-answers', password='pass')
        response = self.client.get(reverse('admin:surveys_surveyanswer_changelist'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '홍길동')
        self.assertContains(response, '팀A')
        self.assertContains(response, '추가로 느낀 증상을 적어주세요')
        self.assertContains(response, '두통과 어지러움이 있습니다.')

    def test_health_submission_detail_uses_snapshot_without_answer_rows(self):
        submission = SurveySubmission.objects.get(team=self.team)
        submission.form_snapshot_json = {
            'sections': [
                {
                    'title': '신체',
                    'questions': [
                        {
                            'key': 'q1',
                            'label': '제출 당시 두통 여부',
                        },
                        {
                            'key': 'q_note',
                            'label': '제출 당시 추가 증상',
                        },
                    ],
                },
            ],
        }
        submission.save(update_fields=['form_snapshot_json', 'updated_at'])
        submission.answers.all().delete()
        self.question.label = '나중에 변경된 두통 문항'
        self.question.save(update_fields=['label', 'updated_at'])
        root = get_user_model().objects.create_superuser(
            username='root-health-detail',
            password='pass',
            email='root-health-detail@example.com',
        )
        self.client.force_login(root)

        response = self.client.get(
            reverse(
                'admin:surveys_surveysubmission_change',
                args=[submission.pk],
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '제출 당시 두통 여부')
        self.assertContains(response, '예')
        self.assertContains(response, '제출 당시 추가 증상')
        self.assertContains(response, '두통과 어지러움이 있습니다.')
        self.assertNotContains(response, '나중에 변경된 두통 문항')

    def test_staff_cannot_open_superuser_only_admin_menus_by_direct_url(self):
        self.client.login(username='staff', password='pass')
        blocked_urls = [
            reverse('admin:teams_respondent_changelist'),
            reverse('admin:surveys_surveyanswer_changelist'),
            reverse('admin:admins_adminteampermission_changelist'),
            reverse('admin:auth_user_changelist'),
            reverse('admin:auth_group_changelist'),
        ]
        for url in blocked_urls:
            response = self.client.get(url)
            self.assertIn(response.status_code, (403, 302))

    def test_superuser_can_delete_respondent_with_related_submission_from_admin(self):
        user_model = get_user_model()
        root = user_model.objects.create_superuser(username='root-delete', password='pass', email='root-delete@example.com')
        respondent = Respondent.objects.get(team=self.team, normalized_name='홍길동')
        submission = SurveySubmission.objects.get(respondent=respondent)

        self.client.force_login(root)
        delete_url = reverse('admin:teams_respondent_delete', args=[respondent.id])

        confirm_response = self.client.get(delete_url)
        self.assertEqual(confirm_response.status_code, 200)
        self.assertContains(confirm_response, respondent.name)

        post_response = self.client.post(delete_url, {'post': 'yes'}, follow=True)
        self.assertEqual(post_response.status_code, 200)
        self.assertFalse(Respondent.objects.filter(id=respondent.id).exists())
        self.assertFalse(SurveySubmission.objects.filter(id=submission.id).exists())


class ExitSurveyExcelAdminActionTest(TestCase):
    def setUp(self):
        self.superuser = get_user_model().objects.create_superuser(
            username='survey-excel-admin',
            password='pass',
            email='survey-excel@example.com',
        )
        self.team = Team.objects.create(
            name='엑셀 테스트 팀',
            organization_name='독도재단',
            organization_name_en='Dokdo Foundation',
            code='EXCEL-TEAM',
            manager_entry_code='EXCEL-MANAGER',
            is_active=True,
            start_date=date.today(),
            max_day_no=1,
        )
        self.form = SurveyForm.objects.create(
            team=self.team,
            day_no=1,
            title='탐방 종료 설문',
            title_en='Post-tour Survey',
            version=1,
            form_type=SurveyForm.FormType.EXIT_SURVEY,
            is_anonymous=False,
            is_active=True,
        )
        section = SurveySection.objects.create(
            form=self.form,
            title='평가',
            title_en='Evaluation',
            section_type=SurveySection.SectionType.SURVEY,
            display_order=1,
        )
        self.question = SurveyQuestion.objects.create(
            section=section,
            key='satisfaction',
            label='전반적으로 만족하셨습니까?',
            label_en='Were you satisfied overall?',
            question_type=SurveyQuestion.QuestionType.SINGLE,
            options_json=[
                {'label': '매우 우수', 'label_en': 'Excellent', 'score': 5},
                {'label': '보통', 'label_en': 'Normal', 'score': 3},
            ],
            is_required=True,
            display_order=1,
        )
        self.text_question = SurveyQuestion.objects.create(
            section=section,
            key='comment',
            label='기타 의견',
            label_en='Other comments',
            question_type=SurveyQuestion.QuestionType.TEXT,
            is_required=False,
            display_order=2,
        )
        snapshot = {
            'title': self.form.title,
            'title_en': self.form.title_en,
            'sections': [
                {
                    'title': section.title,
                    'title_en': section.title_en,
                    'section_type': section.section_type,
                    'questions': [
                        {
                            'id': self.question.id,
                            'key': self.question.key,
                            'label': self.question.label,
                            'label_en': self.question.label_en,
                            'question_type': self.question.question_type,
                            'options': self.question.options_json,
                        },
                        {
                            'id': self.text_question.id,
                            'key': self.text_question.key,
                            'label': self.text_question.label,
                            'label_en': self.text_question.label_en,
                            'question_type': self.text_question.question_type,
                            'options': [],
                        },
                    ],
                },
            ],
        }
        for index, selected in enumerate(('매우 우수', '보통'), start=1):
            respondent = Respondent.objects.create(
                team=self.team,
                name=f'응답자{index}',
                birth_date=f'1990-01-0{index}',
                phone=f'0100000000{index}',
                normalized_name=f'응답자{index}',
                normalized_phone=f'0100000000{index}',
            )
            ExitSurveySubmission.objects.create(
                team=self.team,
                respondent=respondent,
                survey_form=self.form,
                form_version=1,
                form_snapshot_json=snapshot,
                answers_json={
                    'satisfaction': selected,
                    'comment': (
                        '=WEBSERVICE("https://example.invalid")'
                        if index == 1
                        else '좋았습니다.'
                    ),
                },
                submitted_at=timezone.now(),
            )
            Participant.objects.create(
                team=self.team,
                name=f'응답자{index}',
                birth_date=f'1990-01-0{index}',
                phone=f'0100000000{index}',
                normalized_name=f'응답자{index}',
                normalized_phone=f'0100000000{index}',
                is_active=True,
            )
        Participant.objects.create(
            team=self.team,
            name='미제출자',
            birth_date='1990-01-03',
            phone='01000000003',
            normalized_name='미제출자',
            normalized_phone='01000000003',
            is_active=True,
        )

    def test_admin_action_downloads_selected_results_as_aggregate_excel(self):
        self.client.force_login(self.superuser)
        selected_ids = list(
            ExitSurveySubmission.objects.values_list('pk', flat=True)
        )

        response = self.client.post(
            reverse('admin:surveys_exitsurveysubmission_changelist'),
            {
                'action': 'export_survey_results_excel',
                '_selected_action': selected_ids,
                'index': '0',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        audit = AdminDownloadAudit.objects.get(resource_type='exit_survey')
        self.assertEqual(audit.team, self.team)
        self.assertEqual(audit.record_count, 2)
        self.assertEqual(audit.result, 'success')
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ['문항별 집계', '주관식 응답', '내보내기 정보'],
        )
        summary_rows = [
            [cell.value for cell in row]
            for row in workbook['문항별 집계'].iter_rows()
        ]
        headers = summary_rows[1]
        question_row = next(
            row
            for row in summary_rows
            if '전반적으로 만족하셨습니까?' in str(row[1])
        )
        self.assertEqual(
            question_row[headers.index('5점')],
            1,
        )
        self.assertEqual(question_row[headers.index('3점')], 1)
        self.assertEqual(question_row[headers.index('4점')], 0)
        self.assertNotIn('미응답', headers)
        self.assertNotIn('응답자 수', headers)
        self.assertIn('전체 인원 3명', summary_rows[0][0])
        self.assertIn('응답자 2명', summary_rows[0][0])
        info_values = {
            row[0].value: row[1].value
            for row in workbook['내보내기 정보'].iter_rows()
        }
        self.assertEqual(info_values['전체 인원수'], 3)
        self.assertEqual(info_values['응답자 수'], 2)
        self.assertEqual(info_values['주관식 문항 수'], 1)
        self.assertEqual(info_values['주관식 응답 건수'], 2)
        subjective_rows = [
            [cell.value for cell in row]
            for row in workbook['주관식 응답'].iter_rows(min_row=2)
        ]
        self.assertEqual(
            subjective_rows,
            [
                [
                    1,
                    2,
                    '기타 의견',
                    'Other comments',
                    '\'=WEBSERVICE("https://example.invalid")',
                ],
                [1, 2, '기타 의견', 'Other comments', '좋았습니다.'],
            ],
        )
        workbook_values = [
            cell.value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ]
        self.assertIn(
            '\'=WEBSERVICE("https://example.invalid")',
            workbook_values,
        )
        self.assertIn('좋았습니다.', workbook_values)
        self.assertNotIn('응답자1', workbook_values)

    def test_exit_survey_admin_exposes_excel_action(self):
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse('admin:surveys_exitsurveysubmission_changelist')
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            '선택한 설문 제출 결과를 Excel로 다운로드',
        )

    def test_exit_survey_admin_detail_renders_snapshot_answers_not_raw_json(self):
        submission = ExitSurveySubmission.objects.order_by('pk').first()
        self.question.label = '나중에 변경된 문항'
        self.question.save(update_fields=['label', 'updated_at'])
        self.client.force_login(self.superuser)

        response = self.client.get(
            reverse(
                'admin:surveys_exitsurveysubmission_change',
                args=[submission.pk],
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '답변 요약')
        self.assertContains(response, '전반적으로 만족하셨습니까?')
        self.assertContains(response, '매우 우수')
        self.assertContains(response, '기타 의견')
        self.assertContains(
            response,
            '=WEBSERVICE(&quot;https://example.invalid&quot;)',
            html=False,
        )
        self.assertNotContains(response, '나중에 변경된 문항')
        self.assertNotContains(response, 'Answers json')
        self.assertNotContains(response, 'Form snapshot json')
