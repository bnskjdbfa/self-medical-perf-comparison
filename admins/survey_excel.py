from __future__ import annotations

from collections import Counter, OrderedDict
import hashlib
from io import BytesIO
import json
import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from surveys.models import ExitSurveySubmission, SurveyQuestion


INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]+')
EXCEL_FORMULA_PREFIXES = ('=', '+', '-', '@')
SCORE_COLUMNS = (5, 4, 3, 2, 1)


def safe_excel_text(value) -> str:
    text = str(value or '')
    if text.startswith(EXCEL_FORMULA_PREFIXES):
        return f"'{text}"
    return text


def survey_results_filename(*, organization_name: str, team_name: str) -> str:
    download_date = timezone.localdate().isoformat()
    parts = [
        INVALID_FILENAME_CHARS.sub('_', value).strip(' ._')
        for value in (organization_name, team_name)
        if value
    ]
    prefix = '_'.join(parts) or '설문결과'
    return f'{prefix}_설문결과_{download_date}.xlsx'


def _snapshot_questions(snapshot: dict) -> list[dict]:
    questions = []
    for section in (snapshot or {}).get('sections') or []:
        for question in section.get('questions') or []:
            questions.append(question)
    return questions


def _option_entries(question: dict) -> list[dict]:
    entries = []
    for option in question.get('options') or []:
        if isinstance(option, dict):
            label = str(option.get('label') or '').strip()
            raw_score = option.get('score')
        else:
            label = str(option or '').strip()
            raw_score = None
        if label:
            try:
                score = int(raw_score)
            except (TypeError, ValueError):
                score = None
            entries.append({'label': label, 'score': score})
    return entries


def _bilingual_label(korean, english) -> str:
    korean_text = safe_excel_text(korean).strip()
    english_text = safe_excel_text(english).strip()
    if korean_text and english_text:
        return f'{korean_text}\n{english_text}'
    return korean_text or english_text


def _aggregate_submission_group(submissions: list[ExitSurveySubmission]) -> dict:
    snapshot = submissions[0].form_snapshot_json or {}
    questions = _snapshot_questions(snapshot)
    answers = [submission.answers_json or {} for submission in submissions]
    question_rows = []
    subjective_rows = []

    for question_number, question in enumerate(questions, start=1):
        key = str(question.get('key') or '')
        question_type = question.get('question_type')
        question_label_ko = safe_excel_text(question.get('label')).strip()
        question_label_en = safe_excel_text(question.get('label_en')).strip()
        question_label = _bilingual_label(question_label_ko, question_label_en)
        if question_type == SurveyQuestion.QuestionType.TEXT:
            for answer in answers:
                value = str(answer.get(key) or '').strip()
                if value:
                    subjective_rows.append(
                        {
                            'question_number': question_number,
                            'question_ko': question_label_ko,
                            'question_en': question_label_en,
                            'answer': safe_excel_text(value),
                        }
                    )
            continue

        option_entries = _option_entries(question)
        score_by_label = {
            option['label']: option['score']
            for option in option_entries
            if option['score'] in SCORE_COLUMNS
        }
        counts = Counter()
        for answer in answers:
            raw_value = answer.get(key)
            values = raw_value if isinstance(raw_value, list) else [raw_value]
            selected = [str(value).strip() for value in values if str(value or '').strip()]
            if not selected:
                continue
            selected_scores = [
                score_by_label[value]
                for value in selected
                if value in score_by_label
            ]
            if not selected_scores:
                continue
            counts.update(selected_scores)

        question_rows.append(
            {
                'question_number': question_number,
                'question': question_label,
                'counts': counts,
            }
        )

    return {
        'snapshot': snapshot,
        'question_rows': question_rows,
        'subjective_rows': subjective_rows,
        'submission_count': len(submissions),
    }


def _snapshot_fingerprint(snapshot: dict) -> str:
    canonical = json.dumps(
        snapshot or {},
        ensure_ascii=False,
        sort_keys=True,
        separators=(',', ':'),
    )
    return hashlib.sha256(canonical.encode('utf-8')).hexdigest()


def _group_submissions(
    submissions,
) -> list[tuple[tuple[int, int, str], list[ExitSurveySubmission]]]:
    grouped: OrderedDict[
        tuple[int, int, str],
        list[ExitSurveySubmission],
    ] = OrderedDict()
    for submission in submissions:
        key = (
            submission.survey_form_id,
            submission.form_version,
            _snapshot_fingerprint(submission.form_snapshot_json),
        )
        grouped.setdefault(key, []).append(submission)
    return list(grouped.items())


def build_survey_results_workbook(*, team, submissions) -> BytesIO:
    grouped_submissions = _group_submissions(list(submissions))
    total_participant_count = team.participants.filter(is_active=True).count()
    respondent_count = len(
        {
            submission.respondent_id
            for _, group in grouped_submissions
            for submission in group
        }
    )
    workbook = Workbook()
    summary = workbook.active
    summary.title = '문항별 집계'
    subjective = workbook.create_sheet('주관식 응답')
    info = workbook.create_sheet('내보내기 정보')

    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill('solid', fgColor='D9EAF7')
    title_fill = PatternFill('solid', fgColor='4A829E')
    title_font = Font(color='FFFFFF', bold=True, size=13)
    header_font = Font(bold=True)
    centered = Alignment(horizontal='center', vertical='center', wrap_text=True)
    wrapped = Alignment(vertical='center', wrap_text=True)

    summary_row = 1
    subjective.append(
        ['설문 버전', '문항 번호', '국문 문항', '영문 문항', '응답 내용']
    )
    for cell in subjective[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = centered
    subjective_response_count = 0
    subjective_question_keys = set()

    for (
        _form_id,
        form_version,
        _snapshot_hash,
    ), submission_group in grouped_submissions:
        aggregate = _aggregate_submission_group(submission_group)
        snapshot = aggregate['snapshot']
        title = _bilingual_label(
            snapshot.get('title') or '설문 제목 미기록',
            snapshot.get('title_en'),
        )
        last_column = 2 + len(SCORE_COLUMNS)

        summary.merge_cells(
            start_row=summary_row,
            start_column=1,
            end_row=summary_row,
            end_column=last_column,
        )
        title_cell = summary.cell(
            summary_row,
            1,
            safe_excel_text(
                f'{title} · 버전 {form_version}'
                f' · 전체 인원 {total_participant_count}명'
                f' · 응답자 {respondent_count}명'
            ),
        )
        title_cell.fill = title_fill
        title_cell.font = title_font
        title_cell.alignment = wrapped
        summary_row += 1

        headers = [
            '번호',
            '평가 항목',
            *(f'{score}점' for score in SCORE_COLUMNS),
        ]
        for column, header in enumerate(headers, start=1):
            cell = summary.cell(summary_row, column, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = centered
        summary_row += 1

        for row in aggregate['question_rows']:
            values = [row['question_number'], row['question']]
            values.extend(
                row['counts'].get(score, 0)
                for score in SCORE_COLUMNS
            )
            for column, value in enumerate(values, start=1):
                cell = summary.cell(summary_row, column, value)
                cell.border = border
                cell.alignment = wrapped if column == 2 else centered
            summary_row += 1
        summary_row += 2
        for row in aggregate['subjective_rows']:
            subjective.append(
                [
                    form_version,
                    row['question_number'],
                    row['question_ko'],
                    row['question_en'],
                    row['answer'],
                ]
            )
            subjective_response_count += 1
            subjective_question_keys.add(
                (
                    _form_id,
                    form_version,
                    _snapshot_hash,
                    row['question_number'],
                )
            )

    info_rows = [
        ('운영기관', safe_excel_text(team.organization_name)),
        ('팀', safe_excel_text(team.name)),
        ('전체 인원수', total_participant_count),
        ('응답자 수', respondent_count),
        ('선택한 제출 건수', sum(len(group) for _, group in grouped_submissions)),
        ('주관식 문항 수', len(subjective_question_keys)),
        ('주관식 응답 건수', subjective_response_count),
        ('설문 버전 수', len(grouped_submissions)),
        ('다운로드 일시', timezone.localtime().strftime('%Y-%m-%d %H:%M:%S')),
        ('개인정보 포함 여부', '포함하지 않음'),
    ]
    for label, value in info_rows:
        info.append([label, value])
    for row in info.iter_rows():
        row[0].font = header_font
        row[0].fill = header_fill
        for cell in row:
            cell.border = border
            cell.alignment = wrapped
    for row in subjective.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = wrapped

    summary.column_dimensions['A'].width = 8
    summary.column_dimensions['B'].width = 72
    for column in range(3, summary.max_column + 1):
        summary.column_dimensions[get_column_letter(column)].width = 18
    subjective.column_dimensions['A'].width = 12
    subjective.column_dimensions['B'].width = 12
    subjective.column_dimensions['C'].width = 55
    subjective.column_dimensions['D'].width = 55
    subjective.column_dimensions['E'].width = 90
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 50
    summary.freeze_panes = 'A2'
    subjective.freeze_panes = 'A2'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
