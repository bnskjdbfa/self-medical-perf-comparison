"""Bulk roster upload: parse → plan (no writes) → apply (atomic) with an audit row.

Under-age rows (팀 시작일 기준 만 14세 미만) are no longer rejected. They are registered in a
"동의 대기" state that blocks login until an operator records guardian consent, so the roster
can be uploaded once and only the consent checkbox is handled afterwards.
"""
from __future__ import annotations

import csv
import io
import os
from dataclasses import asdict, dataclass
from datetime import date, datetime

from django import forms
from django.db import transaction
from django.utils import timezone

from surveys.services import normalize_name, normalize_phone

from .age_policy import is_eligible_participant_age, participant_has_required_guardian_consent
from .guardian_consent import mark_guardian_consent_verified, record_guardian_consent_event
from .models import GenderChoices, GuardianConsentEvent, Participant, ParticipantUploadAudit, Team

UPLOAD_COLUMN_ALIASES = {
    'name': {'name', '이름'},
    'birth_date': {'birth_date', '생년월일', '주민등록번호', '주민번호', '주민등록번호앞자리'},
    'phone': {'phone', '연락처', '휴대폰번호', '휴대폰'},
    'gender': {'gender', '성별'},
    'guardian_consent_verified': {
        'guardian_consent_verified', '법정대리인동의확인', '보호자동의확인',
    },
}
# 전화번호는 연락처로 저장하되 식별에는 쓰지 않으므로 선택 열이다. 성별은 식별 키의 일부라 필수.
OPTIONAL_UPLOAD_COLUMNS = {'phone', 'guardian_consent_verified'}

ACTION_LABELS = {'create': '생성', 'update': '갱신', 'skip': '건너뜀(이미 등록)', 'error': '오류'}
CONSENT_LABELS = {
    'not_required': '해당 없음',
    'verified': '확인 완료로 기록',
    'pending': '동의 대기 (로그인 불가)',
    'kept': '기존 확인 유지',
}
PENDING_CONSENT_NOTE = '만 14세 미만: 법정대리인 동의 확인 전까지 로그인할 수 없습니다.'


def normalize_upload_header(value: str) -> str:
    return ''.join(str(value or '').strip().lower().split())


def normalize_upload_row(row: dict) -> dict:
    normalized = {}
    normalized_row = {
        normalize_upload_header(key): row[key]
        for key in row.keys()
    }
    for canonical, aliases in UPLOAD_COLUMN_ALIASES.items():
        value = ''
        for alias in aliases:
            normalized_alias = normalize_upload_header(alias)
            if normalized_alias in normalized_row and normalized_row[normalized_alias] is not None:
                value = str(normalized_row[normalized_alias]).strip()
                break
        normalized[canonical] = value
    return normalized


def normalize_upload_phone(value) -> str:
    if value is None:
        return ''
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_upload_gender(value) -> str | None:
    raw = str(value or '').strip().upper()
    mapping = {
        'M': GenderChoices.MALE,
        '남': GenderChoices.MALE,
        '남자': GenderChoices.MALE,
        'MALE': GenderChoices.MALE,
        'F': GenderChoices.FEMALE,
        '여': GenderChoices.FEMALE,
        '여자': GenderChoices.FEMALE,
        'FEMALE': GenderChoices.FEMALE,
        'U': GenderChoices.UNSPECIFIED,
        '미입력': GenderChoices.UNSPECIFIED,
        'UNKNOWN': GenderChoices.UNSPECIFIED,
        '': None,
    }
    if raw not in mapping:
        raise forms.ValidationError('gender/성별 값은 M/F/U 또는 남/여/미입력만 허용됩니다.')
    return mapping[raw]


def normalize_upload_boolean(value) -> bool:
    return str(value or '').strip().lower() in {'1', 'true', 'yes', 'y', '예', '확인', '완료'}


def parse_birth_date(value: str) -> date | None:
    text = str(value or '').strip()
    if text.endswith('.0') and text[:-2].isdigit():
        text = text[:-2]
    raw = ''.join(ch for ch in text if ch.isdigit())
    if len(raw) >= 13:
        raw = raw[:6]
    if len(raw) == 7:
        raw = raw[:6]
    if len(raw) == 5:
        raw = raw.zfill(6)
    try:
        if len(raw) == 8:
            return datetime.strptime(raw, '%Y%m%d').date()
        if len(raw) == 6:
            yy = int(raw[:2])
            prefix = '19' if yy >= 30 else '20'
            return datetime.strptime(f'{prefix}{raw}', '%Y%m%d').date()
    except ValueError:
        return None
    return None


def iter_upload_rows(file_name: str, raw_bytes: bytes):
    ext = os.path.splitext(file_name.lower())[1]
    if ext == '.xlsx':
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover
            raise forms.ValidationError('xlsx 업로드를 위해 openpyxl이 필요합니다. 현재는 csv 업로드를 사용해주세요.') from exc
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell).strip() if cell is not None else '' for cell in rows[0]]
        for row in rows[1:]:
            yield {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        return

    decoded = None
    last_error = None
    for encoding in ('utf-8-sig', 'cp949', 'euc-kr', 'utf-16'):
        try:
            decoded = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    if decoded is None:
        raise forms.ValidationError('CSV 인코딩을 해석할 수 없습니다. UTF-8 또는 CP949(EUC-KR) 형식으로 저장해주세요.') from last_error
    reader = csv.DictReader(io.StringIO(decoded))
    for row in reader:
        yield row


def missing_required_columns(rows: list[dict]) -> list[str]:
    if not rows:
        return []
    first_row_keys = {normalize_upload_header(key) for key in rows[0].keys()}
    return [
        canonical
        for canonical, aliases in UPLOAD_COLUMN_ALIASES.items()
        if canonical not in OPTIONAL_UPLOAD_COLUMNS
        and not first_row_keys.intersection({normalize_upload_header(alias) for alias in aliases})
    ]


@dataclass
class UploadRowPlan:
    row_no: int
    name: str = ''
    birth_date: str = ''
    phone: str = ''
    gender: str | None = None
    action: str = 'error'          # create | update | skip | error
    consent: str = 'not_required'  # not_required | verified | pending | kept
    message: str = ''
    existing_id: int | None = None

    @property
    def is_error(self) -> bool:
        return self.action == 'error'

    @property
    def action_label(self) -> str:
        return ACTION_LABELS.get(self.action, self.action)

    @property
    def consent_label(self) -> str:
        return CONSENT_LABELS.get(self.consent, self.consent)

    def as_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict) -> 'UploadRowPlan':
        return cls(**data)


def plan_participant_upload(*, team: Team, rows: list[dict], update_existing: bool) -> list[UploadRowPlan]:
    """Validate every row and decide what applying would do. Performs no writes."""
    plans: list[UploadRowPlan] = []
    for index, row in enumerate(rows, start=2):
        plan = UploadRowPlan(row_no=index)
        normalized_row = normalize_upload_row(row)
        plan.name = normalized_row['name']
        birth_date = parse_birth_date(normalized_row['birth_date'])
        plan.birth_date = birth_date.isoformat() if birth_date else normalized_row['birth_date']
        plan.phone = normalize_upload_phone(normalized_row['phone'])  # 연락처 정보. 식별에는 쓰지 않으며 비워도 된다.
        try:
            plan.gender = normalize_upload_gender(normalized_row['gender'])
        except forms.ValidationError:
            plan.message = 'gender/성별 값이 올바르지 않습니다.'
            plans.append(plan)
            continue
        if not plan.name or not birth_date:
            plan.message = 'name/birth_date 모두 필요합니다.'
            plans.append(plan)
            continue
        if plan.gender in (None, GenderChoices.UNSPECIFIED):
            plan.message = '성별은 남/여 중 하나여야 합니다 (식별에 사용).'
            plans.append(plan)
            continue

        existing = Participant.objects.filter(
            team=team,
            normalized_name__iexact=normalize_name(plan.name),
            birth_date=birth_date,
            gender=plan.gender,
        ).first()
        is_under_age = not is_eligible_participant_age(birth_date, team.start_date)
        consent_flag = normalize_upload_boolean(normalized_row.get('guardian_consent_verified'))
        if not is_under_age:
            plan.consent = 'not_required'
        elif consent_flag:
            plan.consent = 'verified'
        elif existing and participant_has_required_guardian_consent(existing):
            plan.consent = 'kept'
        else:
            plan.consent = 'pending'

        if existing:
            plan.existing_id = existing.pk
            plan.action = 'update' if update_existing else 'skip'
            if not update_existing:
                plan.message = '이미 등록된 참가자입니다.'
        else:
            plan.action = 'create'
        if plan.consent == 'pending':
            plan.message = f'{plan.message} {PENDING_CONSENT_NOTE}'.strip()
        plans.append(plan)
    return plans


def summarize_plans(plans: list[UploadRowPlan]) -> dict:
    return {
        'total': len(plans),
        'create': sum(p.action == 'create' for p in plans),
        'update': sum(p.action == 'update' for p in plans),
        'skip': sum(p.action == 'skip' for p in plans),
        'error': sum(p.action == 'error' for p in plans),
        'pending': sum(p.consent == 'pending' and p.action in {'create', 'update'} for p in plans),
        'verified': sum(p.consent == 'verified' and p.action in {'create', 'update'} for p in plans),
    }


def apply_participant_upload(
    *,
    team: Team,
    plans: list[UploadRowPlan],
    update_existing: bool,
    actor,
    file_name: str = '',
) -> ParticipantUploadAudit:
    """Apply a plan produced by plan_participant_upload(). All rows succeed or none do."""
    now = timezone.now()
    created = updated = skipped = pending = 0
    with transaction.atomic():
        for plan in plans:
            if plan.action == 'error':
                continue
            if plan.action == 'skip':
                skipped += 1
                continue
            birth_date = date.fromisoformat(plan.birth_date)
            participant = None
            if plan.action == 'update' and plan.existing_id:
                participant = Participant.objects.filter(pk=plan.existing_id, team=team).first()
            if participant is not None:
                participant.name = plan.name
                participant.normalized_name = normalize_name(plan.name)
                participant.gender = plan.gender
                if plan.phone:
                    participant.phone = plan.phone
                    participant.normalized_phone = normalize_phone(plan.phone)
                participant.is_active = True
                newly_verified = plan.consent == 'verified' and not participant.guardian_consent_verified_at
                if newly_verified:
                    mark_guardian_consent_verified(participant, actor=actor, now=now)
                participant.save()
                updated += 1
            else:
                participant = Participant(
                    team=team,
                    name=plan.name,
                    normalized_name=normalize_name(plan.name),
                    birth_date=birth_date,
                    gender=plan.gender,
                    phone=plan.phone,
                    normalized_phone=normalize_phone(plan.phone),
                    is_active=True,
                )
                newly_verified = plan.consent == 'verified'
                if newly_verified:
                    mark_guardian_consent_verified(participant, actor=actor, now=now)
                participant.save()
                created += 1
            if newly_verified:
                record_guardian_consent_event(
                    participant,
                    action=GuardianConsentEvent.Action.VERIFIED,
                    actor=actor,
                    source=GuardianConsentEvent.Source.BULK_UPLOAD,
                    note=f'일괄 등록 파일 {file_name} {plan.row_no}행'.strip(),
                    occurred_at=now,
                )
            if plan.consent == 'pending':
                pending += 1
        return ParticipantUploadAudit.objects.create(
            admin_user=actor if getattr(actor, 'pk', None) else None,
            team=team,
            team_name=team.name,
            file_name=file_name[:255],
            total_rows=len(plans),
            created_count=created,
            updated_count=updated,
            skipped_count=skipped,
            error_count=sum(p.action == 'error' for p in plans),
            pending_consent_count=pending,
            requested_at=now,
        )


def error_rows_csv(plans: list[UploadRowPlan]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['행', '이름', '생년월일', '연락처', '성별', '사유'])
    for plan in plans:
        if plan.is_error:
            writer.writerow([plan.row_no, plan.name, plan.birth_date, plan.phone, plan.gender or '', plan.message])
    return output.getvalue().encode('utf-8-sig')
